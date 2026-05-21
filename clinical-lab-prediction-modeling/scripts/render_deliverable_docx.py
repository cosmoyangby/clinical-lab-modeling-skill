#!/usr/bin/env python3
"""Render a Markdown research decision framework into a Word document."""

from __future__ import annotations

import argparse
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH


def collect_headings(text: str) -> list[str]:
    """Collect second-level headings for a simple manual table of contents."""
    headings: list[str] = []
    for line in text.splitlines():
        match = re.match(r"^##\s+(.*)$", line.strip())
        if match and match.group(1).strip() not in {"Title Page", "Table of Contents"}:
            headings.append(match.group(1).strip())
    return headings


def add_markdown_line(document: Document, line: str) -> None:
    """Add one supported Markdown line to the document."""
    stripped = line.strip()

    if not stripped:
        return

    heading = re.match(r"^(#{1,6})\s+(.*)$", stripped)
    if heading:
        level = min(len(heading.group(1)), 4)
        document.add_heading(heading.group(2).strip(), level=level)
        return

    if stripped.startswith("- "):
        document.add_paragraph(stripped[2:].strip(), style="List Bullet")
        return

    numbered = re.match(r"^\d+\.\s+(.*)$", stripped)
    if numbered:
        document.add_paragraph(numbered.group(1).strip(), style="List Number")
        return

    document.add_paragraph(stripped)


def render_docx(input_md: Path, output_docx: Path) -> None:
    """Render a simple Markdown file to docx."""
    document = Document()
    text = input_md.read_text(encoding="utf-8")
    toc_headings = collect_headings(text)

    in_code_block = False
    code_lines: list[str] = []
    title_written = False
    skip_toc_source_list = False

    for line in text.splitlines():
        stripped = line.strip()

        title = re.match(r"^#\s+(.*)$", stripped)
        if title and not title_written:
            paragraph = document.add_heading(title.group(1).strip(), level=0)
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            document.add_page_break()
            title_written = True
            continue

        if stripped == "## Table of Contents":
            document.add_heading("目录", level=1)
            for heading in toc_headings:
                document.add_paragraph(heading, style="List Number")
            document.add_page_break()
            skip_toc_source_list = True
            continue

        if skip_toc_source_list:
            if not stripped or stripped == "Include a simple manual table of contents for the major sections.":
                continue
            if stripped.startswith("## "):
                skip_toc_source_list = False
            else:
                continue

        if line.strip().startswith("```"):
            if in_code_block:
                if code_lines:
                    document.add_paragraph("\n".join(code_lines))
                code_lines = []
                in_code_block = False
            else:
                in_code_block = True
            continue

        if in_code_block:
            code_lines.append(line)
            continue

        add_markdown_line(document, line)

    output_docx.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_docx)


def main() -> None:
    """Parse arguments and render the Word document."""
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("input_md", type=Path, help="Input Markdown framework file.")
    parser.add_argument("output_docx", type=Path, help="Output Word document path.")
    args = parser.parse_args()

    render_docx(args.input_md, args.output_docx)


if __name__ == "__main__":
    main()
