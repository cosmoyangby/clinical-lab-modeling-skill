"""Build a polished Word research framework for the TBNK landmark simulation."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


OUT = Path(__file__).with_name("research-decision-framework.docx")

BLUE = "1F4E79"
LIGHT_BLUE = "EAF2F8"
PALE_AMBER = "FFF4E0"
LIGHT_GRAY = "F3F5F7"
DARK = "1F2933"
MID = "5B6770"
WHITE = "FFFFFF"


def set_cell_shading(cell, fill: str) -> None:
    """Set table cell background color."""
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color: str = "D8DEE4", size: str = "6") -> None:
    """Set subtle cell borders."""
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=120, start=120, bottom=120, end=120) -> None:
    """Set table cell internal padding in twips."""
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in {"top": top, "start": start, "bottom": bottom, "end": end}.items():
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_width(table, widths: list[float]) -> None:
    """Apply fixed column widths in inches."""
    table.autofit = False
    for row in table.rows:
        for idx, width in enumerate(widths):
            row.cells[idx].width = Inches(width)


def set_font(run, size: float | None = None, bold: bool | None = None, color: str | None = None) -> None:
    """Apply a consistent Chinese/English font stack to a run."""
    run.font.name = "Arial"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def add_para(doc, text: str = "", size: float = 10.5, color: str = DARK, bold: bool = False, align=None):
    """Add a normal paragraph with controlled spacing."""
    para = doc.add_paragraph()
    if align is not None:
        para.alignment = align
    para.paragraph_format.space_after = Pt(6)
    para.paragraph_format.line_spacing = 1.18
    run = para.add_run(text)
    set_font(run, size=size, bold=bold, color=color)
    return para


def add_heading(doc, text: str, level: int = 1) -> None:
    """Add a styled heading."""
    para = doc.add_paragraph()
    para.paragraph_format.space_before = Pt(16 if level == 1 else 10)
    para.paragraph_format.space_after = Pt(7 if level == 1 else 5)
    run = para.add_run(text)
    if level == 1:
        set_font(run, size=18, bold=True, color=BLUE)
    elif level == 2:
        set_font(run, size=13.5, bold=True, color=DARK)
    else:
        set_font(run, size=11.5, bold=True, color=DARK)


def add_callout(doc, title: str, body: str, fill: str = PALE_AMBER) -> None:
    """Add a compact callout box."""
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_width(table, [6.45])
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    set_cell_border(cell, "E2C37A" if fill == PALE_AMBER else "B8D3EA")
    set_cell_margins(cell, 170, 180, 170, 180)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    para = cell.paragraphs[0]
    para.paragraph_format.space_after = Pt(3)
    run = para.add_run(title)
    set_font(run, size=10.5, bold=True, color=DARK)
    para = cell.add_paragraph()
    para.paragraph_format.space_after = Pt(0)
    run = para.add_run(body)
    set_font(run, size=10, color=DARK)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_kv_table(doc, rows: list[tuple[str, str]], widths=(1.75, 4.7), header_fill=LIGHT_BLUE) -> None:
    """Add a key-value table with readable spacing."""
    table = doc.add_table(rows=len(rows), cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_width(table, list(widths))
    for idx, (key, value) in enumerate(rows):
        left, right = table.rows[idx].cells
        set_cell_shading(left, header_fill)
        set_cell_shading(right, WHITE)
        for cell in (left, right):
            set_cell_border(cell)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = left.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(key)
        set_font(r, size=9.5, bold=True, color=BLUE)
        p = right.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(value)
        set_font(r, size=9.5, color=DARK)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def add_matrix(doc, headers: list[str], rows: list[list[str]], widths: list[float]) -> None:
    """Add a compact comparison matrix."""
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_width(table, widths)
    for idx, header in enumerate(headers):
        cell = table.rows[0].cells[idx]
        set_cell_shading(cell, BLUE)
        set_cell_border(cell, BLUE)
        set_cell_margins(cell, 120, 110, 120, 110)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(header)
        set_font(r, size=9, bold=True, color=WHITE)
    for row_idx, values in enumerate(rows, start=1):
        for col_idx, value in enumerate(values):
            cell = table.rows[row_idx].cells[col_idx]
            set_cell_shading(cell, WHITE if row_idx % 2 else "FAFBFC")
            set_cell_border(cell)
            set_cell_margins(cell, 120, 110, 120, 110)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.08
            r = p.add_run(value)
            set_font(r, size=8.8, color=DARK)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def add_bullets(doc, items: list[str]) -> None:
    """Add real Word bullet list."""
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.12
        if p.runs:
            p.runs[0].text = item
            set_font(p.runs[0], size=10, color=DARK)
        else:
            r = p.add_run(item)
            set_font(r, size=10, color=DARK)


def add_numbered(doc, items: list[str]) -> None:
    """Add real Word numbered list."""
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.12
        if p.runs:
            p.runs[0].text = item
            set_font(p.runs[0], size=10, color=DARK)
        else:
            r = p.add_run(item)
            set_font(r, size=10, color=DARK)


def setup_document() -> Document:
    """Create document and base styles."""
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.1)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.15)
    section.right_margin = Cm(2.15)
    section.header_distance = Cm(1.0)
    section.footer_distance = Cm(1.0)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(DARK)

    for style_name in ("List Bullet", "List Number"):
        style = styles[style_name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(10)

    return doc


def add_cover(doc: Document) -> None:
    """Create a clean title page."""
    for _ in range(5):
        doc.add_paragraph()

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("ICU患者TBNK指标预测肺部真菌感染风险")
    set_font(r, size=25, bold=True, color=BLUE)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("48小时landmark预测模型研究框架")
    set_font(r, size=16, bold=False, color=DARK)

    add_para(
        doc,
        "预交付草稿 | 研究设计与分析路线 | 2026-05-28",
        size=10.5,
        color=MID,
        align=WD_ALIGN_PARAGRAPH.CENTER,
    )

    for _ in range(3):
        doc.add_paragraph()

    add_kv_table(
        doc,
        [
            ("研究定位", "48小时landmark风险预测模型开发"),
            ("目标人群", "ICU入院后达到48小时landmark的患者"),
            ("预测窗口", "landmark后14天内肺部真菌PCR阳性"),
            ("当前数据", "约300例；阳性约45例；缺失约5%；暂无外部验证"),
            ("核心提醒", "PCR为临床怀疑后检测，存在结局验证偏倚；死亡/出院导致NA需敏感性分析"),
        ],
        widths=(1.65, 4.8),
    )

    add_callout(
        doc,
        "草稿状态",
        "本文件可用于方法学讨论和后续R agent分析交接；正式分析前仍需确认变量名、PCR检测机制、死亡/出院处理、未检测PCR编码和landmark前阳性病例处理。",
        LIGHT_BLUE,
    )

    doc.add_page_break()


def add_toc(doc: Document) -> None:
    """Create a manual table of contents."""
    add_heading(doc, "目录", 1)
    sections = [
        "一、研究概览",
        "二、数据可行性与研究定位",
        "三、分析方法设计",
        "四、推荐分析流程",
        "五、报告规范与偏倚风险",
        "六、论文方法学草稿",
        "七、R Agent Prompt 文件说明",
        "八、实用注意事项",
    ]
    for item in sections:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.space_after = Pt(5)
        r = p.add_run(item)
        set_font(r, size=11, color=DARK)
    doc.add_page_break()


def build_doc() -> None:
    """Build the polished research framework."""
    doc = setup_document()
    add_cover(doc)
    add_toc(doc)

    add_heading(doc, "一、研究概览", 1)
    add_kv_table(
        doc,
        [
            ("研究问题", "48小时内TBNK指标及早期临床变量能否预测ICU患者随后14天肺部真菌PCR阳性风险。"),
            ("研究类型", "预测模型开发；不是诊断准确性研究。"),
            ("预测起点", "ICU入院后48小时landmark。"),
            ("候选变量", "年龄、性别、CD3、CD4、CD8、NK、B细胞、CD4/CD8、48小时内抗菌药物、48小时内机械通气等。"),
            ("能回答的问题", "能探索早期TBNK模块是否提高14天风险分层能力。"),
            ("不能回答的问题", "不能证明因果关系；不能在无外部验证下声称临床可部署。"),
        ],
    )

    add_heading(doc, "二、数据可行性与研究定位", 1)
    add_callout(
        doc,
        "总体判断：适合探索性或简化预测模型开发",
        "约45个阳性事件可以支持简化模型，但不足以支撑复杂机器学习主模型。建议以惩罚logistic回归或简化logistic回归为主，并采用内部验证。",
    )
    add_matrix(
        doc,
        ["维度", "当前信息", "方法学含义"],
        [
            ["样本量", "约300例", "总体数据量有限，不建议单次随机拆分作为主要验证。"],
            ["阳性事件", "约45例", "候选参数数必须严格控制。"],
            ["缺失值", "约5%", "可处理，但需报告缺失模式并避免插补泄漏。"],
            ["外部验证", "暂无", "只能定位为模型开发或探索性预测模型。"],
            ["结局检测", "临床怀疑后PCR", "存在verification/workup bias。"],
        ],
        [1.3, 1.7, 3.45],
    )

    add_heading(doc, "三、分析方法设计", 1)
    methods = [
        (
            "结局与时间结构",
            "推荐采用48小时landmark设计，仅使用landmark前可获得信息预测后续14天PCR阳性。landmark前已PCR阳性者应排除或单独处理。",
        ),
        (
            "预测变量与特征工程",
            "TBNK指标优先保留连续形式，检查单位、平台、参考区间、极端值和相关性。样本量有限时减少派生变量和交互项。",
        ),
        (
            "缺失值处理",
            "先报告缺失比例和模式。缺失约5%时可比较完整病例与多重插补；插补和标准化必须嵌入重采样流程。",
        ),
        (
            "变量筛选",
            "年龄、性别建议强制纳入；TBNK作为核心模块评估。推荐临床预设少量指标、ridge/elastic net或block-wise比较；避免单因素P值筛选作为主策略。",
        ),
        (
            "建模方法",
            "主模型建议简化logistic或惩罚logistic回归。随机森林、XGBoost等可作探索性比较，但不能作为主要结论。",
        ),
        (
            "验证策略",
            "优先bootstrap内部验证。所有预处理、变量筛选、调参和阈值选择应放入训练或重采样流程内。",
        ),
        (
            "性能评价",
            "报告AUC、PR-AUC、校准曲线、校准截距/斜率、Brier score、阈值表现和决策曲线。不能只报告AUC。",
        ),
        (
            "报告与偏倚控制",
            "按TRIPOD/TRIPOD+AI、PROBAST/PROBAST+AI和BMJ预测模型指南报告，重点解释PCR非系统检测、死亡/出院NA和无外部验证。",
        ),
    ]
    add_matrix(doc, ["模块", "推荐设计"], methods, [1.65, 4.8])

    add_heading(doc, "核心统计学方法速览", 2)
    add_matrix(
        doc,
        ["分析问题", "推荐方法", "主要输出", "注意事项"],
        [
            [
                "描述与基线比较",
                "中位数/IQR或均数/SD；Wilcoxon或t检验；卡方或Fisher精确检验",
                "基线表、阳性/阴性组差异描述",
                "仅用于描述，不作为主变量筛选依据",
            ],
            [
                "缺失值",
                "缺失图谱；Little MCAR检验可选；完整病例与MICE敏感性分析",
                "缺失比例表、插补前后比较",
                "插补应嵌入重采样，避免全数据预插补",
            ],
            [
                "共线性",
                "Pearson/Spearman相关矩阵；VIF；临床相关性聚类",
                "相关热图、VIF表、保留/合并规则",
                "TBNK指标可能高度相关，不机械按阈值删除",
            ],
            [
                "非线性",
                "限制性立方样条；分数多项式可选；似然比检验或AIC辅助判断",
                "非线性效应图、线性/非线性模型比较",
                "事件数有限时只对核心连续变量评估",
            ],
            [
                "变量筛选",
                "临床预设 + ridge/elastic net；bootstrap稳定性；block-wise模块比较",
                "入模变量、选择频率、基础模型与TBNK增量表现",
                "不使用单因素P值筛选作为主策略",
            ],
            [
                "最终建模",
                "简化logistic或惩罚logistic回归；必要时Firth logistic作为敏感性分析",
                "模型公式、OR/系数、预测概率",
                "机器学习仅探索性比较",
            ],
            [
                "模型呈现",
                "列线图；简化风险评分；风险分层表",
                "nomogram、评分规则、低/中/高风险组",
                "只有模型稳定且校准可接受时才推荐",
            ],
            [
                "内部验证",
                "bootstrap乐观校正；重复交叉验证可选",
                "校正AUC、校准斜率、Brier score",
                "不把同源随机拆分称为外部验证",
            ],
            [
                "临床实用性",
                "决策曲线分析；阈值下敏感度/特异度/PPV/NPV",
                "DCA图、阈值表现表",
                "阈值应服务临床决策，不只追求Youden最大",
            ],
        ],
        [1.05, 2.05, 1.75, 1.6],
    )

    add_heading(doc, "变量筛选策略细化", 2)
    add_matrix(
        doc,
        ["项目", "建议"],
        [
            ["候选变量池", "基础临床变量、TBNK核心指标、48小时内临床支持/治疗变量。"],
            ["建模前排除", "48小时后信息、结局组成部分、PCR触发后信息、landmark前已阳性病例。"],
            ["强制纳入", "年龄、性别；TBNK作为核心模块预设评估。"],
            ["推荐路线", "基础临床模型 vs 基础临床+TBNK模块；ridge/elastic net处理共线性。"],
            ["不推荐", "单因素P值筛选、按表观AUC最大化选变量、全部变量普通logistic硬塞。"],
            ["敏感性分析", "bootstrap选择稳定性；不同结局编码下模型表现。"],
        ],
        [1.6, 4.85],
    )

    add_heading(doc, "四、推荐分析流程", 1)
    add_matrix(
        doc,
        ["步骤", "推荐统计学方法", "主要输出"],
        [
            [
                "1. 数据与时间轴审计",
                "landmark规则检查；变量测量时间核对；PCR/死亡/出院时间顺序审计",
                "可纳入样本、排除样本、潜在泄漏变量清单",
            ],
            [
                "2. 结局构建",
                "14天二分类结局；死亡/出院/未检测PCR多规则敏感性编码",
                "主结局变量、敏感性结局变量、结局流程图",
            ],
            [
                "3. 描述与缺失处理",
                "描述统计；组间描述性检验；缺失图谱；完整病例与MICE敏感性分析",
                "Table 1、缺失表、插补策略说明",
            ],
            [
                "4. TBNK数据质量与特征工程",
                "异常值检查；单位/平台核对；连续变量转换；核心变量样条探索",
                "TBNK分布图、异常值规则、变换规则",
            ],
            [
                "5. 共线性与变量筛选",
                "相关矩阵；VIF；临床聚类；ridge/elastic net；bootstrap稳定性选择",
                "相关热图、VIF表、入模变量、选择频率",
            ],
            [
                "6. 模型开发",
                "基础临床logistic；基础临床+TBNK；惩罚logistic主模型",
                "模型公式、系数/OR、预测概率、TBNK增量价值",
            ],
            [
                "7. 模型呈现",
                "列线图；风险评分；风险分层；校准后概率表",
                "nomogram、评分表、低/中/高风险分层",
            ],
            [
                "8. 验证与评价",
                "bootstrap内部验证；AUC/PR-AUC；校准曲线；Brier score；DCA",
                "性能表、校准图、ROC/PR曲线、决策曲线",
            ],
        ],
        [1.35, 3.1, 2.0],
    )

    add_heading(doc, "五、报告规范与偏倚风险", 1)
    add_callout(
        doc,
        "最关键偏倚",
        "PCR为临床怀疑后检测，模型可能预测的是“被怀疑并PCR阳性”的风险，而不是系统筛查下的真实感染发生风险。",
        PALE_AMBER,
    )
    add_bullets(
        doc,
        [
            "学术表述应为48小时landmark风险预测模型开发，不应写成诊断准确性研究。",
            "死亡/出院导致结局NA不能简单当普通缺失，建议进行敏感性分析。",
            "无外部验证时，结论应保持探索性，不宣称临床可用。",
            "报告应覆盖TRIPOD/TRIPOD+AI、PROBAST/PROBAST+AI和BMJ预测模型开发与验证要点。",
        ]
    )

    add_heading(doc, "六、论文方法学草稿", 1)
    add_heading(doc, "中文版本", 2)
    add_para(
        doc,
        "本研究拟开发一个基于ICU入院后48小时内可获得信息的14天肺部真菌感染风险预测模型。研究对象为ICU入院患者，预测起点设定为入院后48小时。候选预测变量包括年龄、性别、48小时内TBNK细胞亚群指标（CD3、CD4、CD8、NK、B细胞及CD4/CD8比值）以及48小时内抗菌药物使用和机械通气等临床变量。研究结局定义为预测起点后14天内临床怀疑后进行肺部真菌PCR检测并呈阳性。",
    )
    add_para(
        doc,
        "考虑到阳性事件数约45例，主模型拟采用简化logistic回归或惩罚logistic回归，并优先考虑ridge或elastic net以降低过拟合和共线性影响。模型性能将通过bootstrap内部验证进行评估，并报告区分度、校准度、Brier score、PR-AUC、阈值相关指标和决策曲线分析。",
    )

    add_heading(doc, "English Version", 2)
    add_para(
        doc,
        "This study aims to develop a 48-hour landmark prediction model for 14-day pulmonary fungal infection risk among ICU patients. Candidate predictors will include age, sex, TBNK lymphocyte subset measurements obtained within the first 48 hours, and clinical variables available within the same time window. The outcome will be defined as pulmonary fungal PCR positivity within 14 days after the landmark time among patients tested because of clinical suspicion.",
    )
    add_para(
        doc,
        "Given approximately 45 positive events, the primary model will use a parsimonious logistic regression or penalized logistic regression approach, with ridge regression or elastic net considered to reduce overfitting and collinearity. Model performance will be assessed using bootstrap internal validation and reported with discrimination, calibration, prediction error, threshold-specific performance, and clinical utility measures.",
    )

    add_heading(doc, "七、R Agent Prompt 文件说明", 1)
    add_kv_table(
        doc,
        [
            ("文件用途", "交给R代码agent执行数据检查、模型开发、内部验证和结果输出。"),
            ("使用前提", "需补充数据路径、变量名、结局编码、PCR检测时间、死亡/出院时间和landmark前阳性处理规则。"),
            ("文件路径", "deliverables/tbnk-fungal-48h-landmark/r-agent-prompt.md"),
        ],
    )

    add_heading(doc, "八、实用注意事项", 1)
    add_bullets(
        doc,
        [
            "当前文件为预交付草稿，不是最终统计分析方案。",
            "48小时后出现的变量不得进入主模型。",
            "阳性事件约45例，模型复杂度必须控制。",
            "AUC不能单独作为模型有效证据，必须同时报告校准和临床实用性。",
            "未获得外部验证前，模型应定位为探索性或开发阶段模型。",
        ]
    )

    footer = doc.sections[0].footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("TBNK 48h Landmark Prediction Framework | Draft v0.1")
    set_font(run, size=8.5, color=MID)

    doc.save(OUT)


if __name__ == "__main__":
    build_doc()
